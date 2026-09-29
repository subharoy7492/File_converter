from __future__ import annotations

import hashlib
import logging
import os
import shutil
import tempfile
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

import streamlit as st
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from converters import CONVERTERS, get_converter


logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)


# --- Page Configuration ---
st.set_page_config(
    page_title="ANSI to Unicode Document Converter",
    layout="centered",
    initial_sidebar_state="collapsed",
)
/* 1. Hide the top-right header menu (includes GitHub icon, Deploy button, and status) */
header[data-testid="stHeader"] {
    display: none !important;
}

/* 2. Reduce top padding of the main container */
.block-container {
    padding-top: 0.5rem !important; /* Adjust down from standard 6rem to lower/raise text */
}

# --- Tailwind CSS-Inspired Modern Theme Injections ---
st.markdown(
    """
    <style>
    /* Google Fonts */
    @import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700&display=swap');

    html, body, [class*="css"] {
        font-family: 'Inter', sans-serif;
    }

    /* Main Container Padding */
    .block-container {
        padding-top: 2rem;
        padding-bottom: 3rem;
        max-width: 720px;
    }

    /* Header Styling */
    .app-header {
        text-align: center;
        margin-bottom: 2rem;
    }

    .app-title {
        font-size: 2.25rem;
        font-weight: 800;
        color: #0f172a; /* Tailwind Slate 900 */
        letter-spacing: -0.025em;
        margin-bottom: 0.5rem;
        display: flex;
        align-items: center;
        justify-content: center;
        gap: 0.5rem;
    }

    .app-subtitle {
        font-size: 1rem;
        color: #64748b; /* Tailwind Slate 500 */
        font-weight: 400;
    }

    /* Card Wrapper */
    div[data-testid="stForm"] {
        background-color: #ffffff;
        border: 1px solid #e2e8f0; /* Tailwind Slate 200 */
        border-radius: 12px;
        padding: 2rem;
        box-shadow: 0 4px 6px -1px rgba(0, 0, 0, 0.05), 0 2px 4px -1px rgba(0, 0, 0, 0.03);
    }

    /* Input Field Labels */
    label[data-testid="stWidgetLabel"] p {
        font-weight: 600 !important;
        color: #334155 !important; /* Tailwind Slate 700 */
        font-size: 0.875rem !important;
    }

    /* Primary Submit Button (Tailwind Blue) */
    div[data-testid="stFormSubmitButton"] button {
        background-color: #2563eb !important; /* Tailwind Blue 600 */
        color: #ffffff !important;
        font-weight: 600 !important;
        border-radius: 8px !important;
        border: none !important;
        padding: 0.625rem 1.25rem !important;
        box-shadow: 0 1px 2px 0 rgba(0, 0, 0, 0.05) !important;
        transition: all 0.2s ease-in-out !important;
    }

    div[data-testid="stFormSubmitButton"] button:hover {
        background-color: #1d4ed8 !important; /* Tailwind Blue 700 */
        box-shadow: 0 4px 6px -1px rgba(37, 99, 235, 0.2) !important;
    }

    /* Download Button (Tailwind Emerald) */
    div[data-testid="stDownloadButton"] button {
        background-color: #059669 !important; /* Tailwind Emerald 600 */
        color: #ffffff !important;
        font-weight: 600 !important;
        border-radius: 8px !important;
        border: none !important;
        padding: 0.625rem 1.25rem !important;
        box-shadow: 0 1px 2px 0 rgba(0, 0, 0, 0.05) !important;
        transition: all 0.2s ease-in-out !important;
    }

    div[data-testid="stDownloadButton"] button:hover {
        background-color: #047857 !important; /* Tailwind Emerald 700 */
        box-shadow: 0 4px 6px -1px rgba(5, 150, 105, 0.2) !important;
    }

    /* File Uploader Custom Styling */
    div[data-testid="stFileUploader"] {
        border-radius: 8px;
    }

    /* Footer */
    .app-footer {
        text-align: center;
        margin-top: 2.5rem;
        padding-top: 1.5rem;
        border-top: 1px solid #f1f5f9; /* Tailwind Slate 100 */
        font-size: 0.875rem;
        color: #94a3b8; /* Tailwind Slate 400 */
    }

    .app-footer span {
        font-weight: 600;
        color: #64748b;
    }
    </style>
    """,
    unsafe_allow_html=True,
)

ANSI_FONT_KEYWORDS = [
    "sutonny",
    "bijoy",
    "stm",
    "soumili",
    "ansi",
    "bangla",
    "bengali",
]

SUPPORTED_EXTENSIONS = ["docx", "idml"]


def is_ansi_font(font_name: str | None) -> bool:
    if not font_name:
        return True
    font_lower = str(font_name).lower()
    return any(keyword in font_lower for keyword in ANSI_FONT_KEYWORDS)


def get_run_font_name(run) -> str | None:
    font_name = run.font.name
    if font_name:
        return font_name

    r_pr = run._element.rPr
    if r_pr is None or r_pr.rFonts is None:
        return None

    r_fonts = r_pr.rFonts
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        value = r_fonts.get(qn(f"w:{attribute}"))
        if value:
            return value

    return None


def set_run_font(run, target_font: str) -> None:
    run.font.name = target_font
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts

    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)

    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{attribute}"), target_font)


def process_runs(runs, target_font: str, converter_instance) -> None:
    for run in runs:
        if not run.text or not run.text.strip():
            continue

        font_name = get_run_font_name(run)

        if is_ansi_font(font_name):
            run.text = converter_instance.convert(run.text)
            set_run_font(run, target_font)


def process_table(table, target_font: str, converter_instance) -> None:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                process_runs(
                    paragraph.runs,
                    target_font,
                    converter_instance,
                )

            for nested_table in cell.tables:
                process_table(
                    nested_table,
                    target_font,
                    converter_instance,
                )


def process_docx(
    input_path: str,
    output_path: str,
    target_font: str,
    converter_instance,
) -> None:
    doc = Document(input_path)

    for paragraph in doc.paragraphs:
        process_runs(
            paragraph.runs,
            target_font,
            converter_instance,
        )

    for table in doc.tables:
        process_table(
            table,
            target_font,
            converter_instance,
        )

    doc.save(output_path)


def safe_extract_zip(zip_file: zipfile.ZipFile, destination: str) -> None:
    destination_path = Path(destination).resolve()

    for member in zip_file.infolist():
        member_path = Path(member.filename)

        if member_path.is_absolute():
            raise ValueError("Unsafe IDML archive path detected.")

        target_path = (destination_path / member_path).resolve()

        try:
            common_path = os.path.commonpath(
                [str(destination_path), str(target_path)]
            )
        except ValueError as exc:
            raise ValueError("Unsafe IDML archive path detected.") from exc

        if common_path != str(destination_path):
            raise ValueError("Unsafe IDML archive path detected.")

        if member.is_dir():
            target_path.mkdir(parents=True, exist_ok=True)
            continue

        target_path.parent.mkdir(parents=True, exist_ok=True)

        with zip_file.open(member) as source, open(target_path, "wb") as target:
            shutil.copyfileobj(source, target)


def process_idml(
    input_path: str,
    output_path: str,
    target_font: str,
    converter_instance,
) -> None:
    extract_dir = f"{input_path}_extracted"

    if os.path.exists(extract_dir):
        shutil.rmtree(extract_dir)

    os.makedirs(extract_dir, exist_ok=True)

    try:
        with zipfile.ZipFile(input_path, "r") as zip_ref:
            safe_extract_zip(zip_ref, extract_dir)

        stories_dir = os.path.join(extract_dir, "Stories")

        if os.path.exists(stories_dir):
            for story_file in os.listdir(stories_dir):
                if not story_file.endswith(".xml"):
                    continue

                file_path = os.path.join(stories_dir, story_file)

                try:
                    tree = ET.parse(file_path)
                    root = tree.getroot()
                    modified = False

                    for element in root.iter():
                        applied_font = element.attrib.get("AppliedFont", "")

                        if not is_ansi_font(applied_font):
                            continue

                        if element.tag.endswith("Content") and element.text:
                            element.text = converter_instance.convert(
                                element.text
                            )
                            modified = True

                        if "AppliedFont" in element.attrib:
                            element.attrib["AppliedFont"] = target_font
                            modified = True

                    if modified:
                        tree.write(
                            file_path,
                            encoding="utf-8",
                            xml_declaration=True,
                        )

                except ET.ParseError:
                    logger.warning(
                        "Skipping invalid XML file: %s",
                        file_path,
                    )

        files_to_write = []

        for folder_name, _, filenames in os.walk(extract_dir):
            for filename in filenames:
                file_path = os.path.join(folder_name, filename)
                archive_name = os.path.relpath(
                    file_path,
                    extract_dir,
                ).replace(os.sep, "/")

                files_to_write.append((file_path, archive_name))

        files_to_write.sort(
            key=lambda item: (
                item[1] != "mimetype",
                item[1],
            )
        )

        with zipfile.ZipFile(
            output_path,
            "w",
            zipfile.ZIP_DEFLATED,
        ) as zip_out:
            for file_path, archive_name in files_to_write:
                if archive_name == "mimetype":
                    zip_out.write(
                        file_path,
                        archive_name,
                        compress_type=zipfile.ZIP_STORED,
                    )
                else:
                    zip_out.write(file_path, archive_name)

    finally:
        shutil.rmtree(extract_dir, ignore_errors=True)


def get_converter_name(key: str, converter_info) -> str:
    if hasattr(converter_info, "name"):
        return converter_info.name

    if isinstance(converter_info, dict):
        return converter_info.get("name", key)

    return key


def get_converter_options():
    keys = list(CONVERTERS.keys())

    names = {
        key: get_converter_name(key, CONVERTERS[key])
        for key in keys
    }

    return keys, names


def clear_previous_result() -> None:
    st.session_state.pop("conversion_result", None)


# --- Header Section ---
st.markdown(
    """
    <div class="app-header">
        <div class="app-title">
            ANSI Document Converter
        </div>
        <p class="app-subtitle">Effortlessly convert legacy Bijoy, STM & Soumili text to Unicode</p>
    </div>
    """,
    unsafe_allow_html=True,
)


converter_keys, converter_names = get_converter_options()

if not converter_keys:
    st.error("No converter engines were found in converters.py.")
    st.stop()


default_converter_index = (
    converter_keys.index("bijoy_v1")
    if "bijoy_v1" in converter_keys
    else 0
)


# --- Input Form ---
with st.form("conversion_form"):
    uploaded_file = st.file_uploader(
        "Upload Document",
        type=SUPPORTED_EXTENSIONS,
        help="Select a .docx or .idml file to process",
    )

    col1, col2 = st.columns(2)

    with col1:
        converter_key = st.selectbox(
            "Converter Engine",
            options=converter_keys,
            index=default_converter_index,
            format_func=lambda key: converter_names[key],
        )

    with col2:
        target_font = st.selectbox(
            "Target Font",
            options=["Kalpurush", "Vrinda"],
            index=0,
        )

    submitted = st.form_submit_button(
        "Convert Document",
        use_container_width=True,
    )


current_signature = None
file_bytes = None

if uploaded_file is not None:
    file_bytes = uploaded_file.getvalue()

    current_signature = (
        hashlib.sha256(file_bytes).hexdigest(),
        converter_key,
        target_font,
    )

    previous_signature = st.session_state.get("source_signature")

    if previous_signature != current_signature:
        clear_previous_result()
        st.session_state["source_signature"] = current_signature
else:
    clear_previous_result()
    st.session_state.pop("source_signature", None)


# --- Processing Logic ---
if submitted:
    if uploaded_file is None or file_bytes is None:
        st.error("Please select a valid DOCX or IDML file first.")
        st.stop()

    original_filename = Path(uploaded_file.name).name
    extension = Path(original_filename).suffix.lower()

    if extension not in {".docx", ".idml"}:
        st.error("Unsupported file format.")
        st.stop()

    try:
        converter_instance = get_converter(converter_key)
    except ValueError as exc:
        st.error(str(exc))
        st.stop()

    with st.spinner("Converting document..."):
        try:
            with tempfile.TemporaryDirectory() as temp_dir:
                input_path = os.path.join(
                    temp_dir,
                    f"input{extension}",
                )

                output_filename = f"Unicode_{original_filename}"
                output_path = os.path.join(
                    temp_dir,
                    output_filename,
                )

                with open(input_path, "wb") as input_file:
                    input_file.write(file_bytes)

                if extension == ".docx":
                    process_docx(
                        input_path,
                        output_path,
                        target_font,
                        converter_instance,
                    )
                    mime_type = (
                        "application/vnd.openxmlformats-"
                        "officedocument.wordprocessingml.document"
                    )

                elif extension == ".idml":
                    process_idml(
                        input_path,
                        output_path,
                        target_font,
                        converter_instance,
                    )
                    mime_type = "application/vnd.adobe.indesign-idml-package"

                with open(output_path, "rb") as output_file:
                    converted_bytes = output_file.read()

            st.session_state["conversion_result"] = {
                "data": converted_bytes,
                "filename": output_filename,
                "mime": mime_type,
            }

            st.toast("Conversion completed!", icon="🎉")

        except Exception as exc:
            logger.exception("Document conversion failed")
            st.error(f"Conversion failed: {exc}")


# --- Output / Download Action ---
conversion_result = st.session_state.get("conversion_result")

if conversion_result:
    st.write("")
    st.download_button(
        label="Download Converted Document",
        data=conversion_result["data"],
        file_name=conversion_result["filename"],
        mime=conversion_result["mime"],
        use_container_width=True,
    )


# --- Footer ---
st.markdown(
    """
    <div class="app-footer">
        Powered by <span>Bangla Haraf Font Foundry</span>
    </div>
    """,
    unsafe_allow_html=True,
)
